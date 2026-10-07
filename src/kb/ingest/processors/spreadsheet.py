"""
Spreadsheet processor: Excel workbooks (`.xlsx`, `.xlsm`) and delimited text (`.csv`,
`.tsv`) converted to markdown tables, one `## <sheet>` section per worksheet.

Excel values are read with `data_only=True`, i.e. the last cached formula results. Legacy
`.xls` isn't supported (it needs a different reader). Workbooks are converted, so
`retain_original` is on for them; the walker keeps the raw bytes when a blob store is set.
"""

import csv
from pathlib import Path

from kb.ingest.processors.base import ProcessedDocument


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def rows_to_markdown(rows: list[list]) -> str:
    """A markdown table; the first row is the header. Empty rows are dropped."""
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    cells = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in cells[1:]]
    return "\n".join(lines)


class SpreadsheetProcessor:
    name = "spreadsheet"
    extensions = frozenset({".xlsx", ".xlsm", ".csv", ".tsv"})

    # The walker reads this class-level flag per processor, so csv/tsv keep their
    # originals too (harmless: small, and deduped by checksum).
    retain_original = True

    def process(self, path: Path) -> ProcessedDocument:
        suffix = path.suffix.lower()
        if suffix in {".csv", ".tsv"}:
            sections = [("", self._read_delimited(path, "\t" if suffix == ".tsv" else ","))]
        else:
            sections = self._read_workbook(path)

        parts = []
        for sheet, rows in sections:
            table = rows_to_markdown(rows)
            if table:
                parts.append(f"## {sheet}\n\n{table}" if sheet else table)
        if not parts:
            raise ValueError("no data found in spreadsheet")
        body = "\n\n".join(parts)
        return ProcessedDocument(title=path.stem, content=f"# {path.stem}\n\n{body}\n")

    @staticmethod
    def _read_delimited(path: Path, delimiter: str) -> list[list]:
        with path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f, delimiter=delimiter))

    @staticmethod
    def _read_workbook(path: Path) -> list[tuple[str, list[list]]]:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
